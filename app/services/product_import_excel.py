"""Excel template + intelligent column mapping for bulk product import."""

from __future__ import annotations

import io
import re
import secrets
from datetime import UTC, datetime
from typing import Any

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

TEMPLATE_VERSION = "DukaPlus_Inventory_Import_v2_rfq"

# Canonical field → header labels (EN + SW + common retail/pharmacy/spare parts)
FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "name": (
        "product_name",
        "product name",
        "name",
        "item",
        "item name",
        "description",
        "bidhaa",
        "jina",
        "jina la bidhaa",
        "part name",
        "drug name",
        "medicine",
    ),
    "sku": ("sku", "code", "item code", "product code", "part no", "part number", "namba", "sku code"),
    "barcode": ("barcode", "bar code", "ean", "upc", "msimbo", "scan code"),
    "category": ("category", "main category", "type", "group", "aina", "kategoria"),
    "unit": ("unit", "uom", "measure", "kipimo", "units"),
    "cost": (
        "cost",
        "buy price",
        "purchase price",
        "cost_buy_price",
        "bei ya kununua",
        "unit cost",
        "landing cost",
    ),
    "price": (
        "price",
        "sell price",
        "selling price",
        "price_sell",
        "retail price",
        "bei ya kuuza",
        "sale price",
    ),
    "stock": ("stock", "stock_qty", "quantity", "qty", "on hand", "idadi", "stoo", "balance"),
    "reorder_point": ("reorder", "reorder_level", "reorder point", "min stock", "kiwango cha chini"),
    "supplier": ("supplier", "vendor", "msambazaji", "supplier name", "manufacturer"),
    "batch_number": ("batch", "batch_number", "batch no", "lot", "lot number"),
    "expiry_date": ("expiry", "expiry_date", "exp date", "expire date", "tarehe ya kuisha", "best before"),
    "requires_prescription": ("rx", "requires_prescription", "prescription", "dawa ya daktari", "requires rx"),
    "location": ("location", "shelf", "bin", "rack", "mahali", "shelf_location"),
    "description": ("notes", "note", "long description", "maelezo", "specification", "spec"),
    "vat_type": ("vat", "vat_type", "tax type", "vat type"),
    "cost_includes_vat": (
        "cost includes vat",
        "cost_includes_vat",
        "price includes vat",
        "incl vat",
        "including vat",
        "bei pamoja na vat",
        "pamoja na vat",
    ),
    "purchase_vat_rate": ("purchase vat", "purchase_vat_rate", "vat rate", "tax rate"),
    "line_tax": (
        "line tax",
        "line_tax",
        "purchase tax",
        "tax id",
        "kodi",
        "kodi ya ununuzi",
        "purchase vat line",
    ),
}

# PO / RFQ line columns (Bidhaa | Idadi | Bei ya Kununua | Bei ya Kuuza | Kodi …)
RFQ_LINE_COLUMNS: list[tuple[str, str, str]] = [
    ("bidhaa", "Required — product name (Bidhaa)", "Activated Charcoal 250mg x20"),
    ("sku", "Required — unique code", "MED-CHAR-250"),
    ("idadi", "Required — quantity (Idadi)", "48"),
    ("bei_ya_kununua", "Required — buy price TZS ex-VAT (Bei ya Kununua)", "1800"),
    ("bei_ya_kuuza", "Required — sell price TZS (Bei ya Kuuza)", "2500"),
    ("kodi", "Per line: none | vat_18 (same as PO)", "none"),
    ("kipimo", "Unit (pcs, box…)", "box"),
    ("msambazaji", "Supplier name (or use Msambazaji sheet)", "Keko Pharma Ltd"),
    ("barcode", "Optional scan code", ""),
    ("category", "Category", "Medicine"),
    ("batch_number", "Batch / lot", "BATCH-2026-01"),
    ("expiry_date", "YYYY-MM-DD", "2027-12-31"),
    ("vat_type", "Product class: standard | exempt | zero", "standard"),
    ("reorder_level", "Low-stock alert", "10"),
    ("shelf_location", "Shelf / bin", "Aisle 3"),
    ("notes", "Extra line notes", ""),
]

SUPPLIER_SHEET_ROWS: list[tuple[str, str, str, str]] = [
    ("supplier_name", "Msambazaji / Supplier name", "Keko Pharma Ltd", "required"),
    ("contact_person", "Mhusika / Contact person", "Sales Desk", "optional"),
    ("phone", "Simu / Phone", "+255712000000", "recommended"),
    ("email", "Barua pepe / Email", "sales@vendor.co.tz", "optional"),
    ("category", "Aina / Category", "Pharmacy supplies", "optional"),
    ("payment_terms", "Masharti ya malipo / Payment terms", "Net 30 days", "optional"),
    ("address", "Anwani / Address", "Dar es Salaam", "optional"),
    ("tin", "TIN / TRA number", "", "optional"),
    ("expected_date", "Tarehe inayotarajiwa / Expected date", "2026-10-07", "optional"),
    ("purchase_vat_scope", "Order VAT: none | all | vat_products", "none", "optional"),
    ("vat_note", "Kumbuka VAT / VAT note", "Bei bila VAT; VAT 18% inaongezwa", "optional"),
]

DEFAULT_PURCHASE_VAT_RATE = 0.18


def _vat_type_is_taxable(vat_type: str | None) -> bool:
    if not vat_type:
        return False
    v = vat_type.lower().strip()
    return v in ("standard", "vat", "vat_18", "taxable", "vatable")


def resolve_unit_cost_ex_vat(
    raw_cost: float,
    *,
    vat_type: str | None,
    cost_includes_vat: bool | None,
    purchase_vat_rate: float | None,
) -> tuple[float, list[str]]:
    """Match PO lines: stored inventory cost is ex-VAT unit cost (TZS)."""
    messages: list[str] = []
    cost = float(raw_cost or 0)
    if cost <= 0:
        return 0.0, messages

    incl = cost_includes_vat
    if incl is None and _vat_type_is_taxable(vat_type):
        incl = False

    rate = purchase_vat_rate if purchase_vat_rate is not None else DEFAULT_PURCHASE_VAT_RATE
    if rate > 1:
        rate = rate / 100.0
    rate = max(0.0, min(rate, 1.0))

    if incl:
        ex = round(cost / (1.0 + rate), 2) if rate > 0 else round(cost, 2)
        messages.append(f"Cost converted from VAT-inclusive to ex-VAT ({int(rate * 100)}%)")
        return ex, messages
    return round(cost, 2), messages


def weighted_average_unit_cost(
    prev_stock: float,
    prev_cost: float,
    add_qty: float,
    incoming_unit_cost: float,
) -> float:
    """Same formula as PO receive in SuppliersView."""
    if add_qty <= 0:
        return round(prev_cost, 2) if prev_cost > 0 else round(incoming_unit_cost, 2)
    new_stock = prev_stock + add_qty
    if new_stock <= 0:
        return round(incoming_unit_cost, 2)
    if prev_stock <= 0:
        return round(incoming_unit_cost, 2)
    blended = ((prev_stock * prev_cost) + (add_qty * incoming_unit_cost)) / new_stock
    return round(max(blended, 0.01), 2)


def resolve_inventory_cost_after_import(
    *,
    prev_stock: float,
    prev_cost: float,
    target_stock: float,
    unit_cost_ex_vat: float,
    stock_mode: str = "absolute",
) -> tuple[float, float, float]:
    """
    Returns (product_cost, stock_delta_for_movement, new_stock).
    stock_mode absolute: Excel stock_qty is on-hand count (opening inventory).
    stock_mode add: Excel stock_qty is quantity received (like PO).
    """
    prev_stock = max(0.0, float(prev_stock or 0))
    prev_cost = max(0.0, float(prev_cost or 0))
    qty_file = max(0.0, float(target_stock or 0))
    unit = max(0.0, float(unit_cost_ex_vat or 0))

    if stock_mode == "add":
        add_qty = qty_file
        new_stock = prev_stock + add_qty
        if add_qty > 0:
            cost = weighted_average_unit_cost(prev_stock, prev_cost, add_qty, unit)
        else:
            cost = prev_cost if prev_cost > 0 else unit
        return cost, add_qty, new_stock

    # absolute (opening / recount)
    new_stock = qty_file
    delta = new_stock - prev_stock
    if delta > 0 and prev_stock > 0 and unit > 0:
        cost = weighted_average_unit_cost(prev_stock, prev_cost, delta, unit)
    elif new_stock > 0 and unit > 0:
        cost = round(unit, 2)
    elif unit > 0:
        cost = round(unit, 2)
    else:
        cost = prev_cost
    move_qty = delta if delta > 0 else (new_stock if prev_stock <= 0 and new_stock > 0 else 0.0)
    return cost, move_qty, new_stock

# Legacy v1 flat columns (still auto-detected on custom sheets)
TEMPLATE_COLUMNS: list[tuple[str, str, str]] = [
    ("product_name", "Required — product label", "Paracetamol 500mg"),
    ("sku", "Required — unique code", "MED-PARA-500"),
    ("cost_buy_price", "Buy price ex-VAT", "1200"),
    ("price_sell", "Sell price", "1500"),
    ("stock_qty", "Quantity", "48"),
    ("supplier", "Vendor", "General Supplier"),
]


def _parse_line_tax(val: Any) -> str:
    s = str(val or "").strip().lower().replace(" ", "_")
    if s in ("vat_18", "vat18", "vat", "18", "18%", "standard", "yes", "ndio"):
        return "vat_18"
    return "none"


def _line_tax_applies_vat(line_tax: str) -> bool:
    return _parse_line_tax(line_tax) == "vat_18"


def _norm_header(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (h or "").strip().lower()).strip()


def _map_headers(header_cells: list[str]) -> dict[int, str]:
    """Map column index → canonical field name."""
    mapping: dict[int, str] = {}
    for idx, raw in enumerate(header_cells):
        norm = _norm_header(str(raw or ""))
        if not norm:
            continue
        for field, aliases in FIELD_ALIASES.items():
            if norm in aliases or norm.replace(" ", "_") in aliases:
                if field not in mapping.values():
                    mapping[idx] = field
                break
        # Direct match to template column names
        template_map = {
            "product name": "name",
            "product_name": "name",
            "bidhaa": "name",
            "cost buy price": "cost",
            "cost_buy_price": "cost",
            "bei ya kununua": "cost",
            "bei_ya_kununua": "cost",
            "price sell": "price",
            "price_sell": "price",
            "bei ya kuuza": "price",
            "bei_ya_kuuza": "price",
            "stock qty": "stock",
            "stock_qty": "stock",
            "idadi": "stock",
            "reorder level": "reorder_point",
            "reorder_level": "reorder_point",
            "shelf location": "location",
            "shelf_location": "location",
            "kipimo": "unit",
            "msambazaji": "supplier",
            "kodi": "line_tax",
            "notes": "description",
        }
        if norm in template_map and template_map[norm] not in mapping.values():
            mapping[idx] = template_map[norm]
    return mapping


def _parse_bool(val: Any) -> bool:
    s = str(val or "").strip().lower()
    return s in ("yes", "y", "true", "1", "ndio", "rx", "required")


def _parse_float(val: Any, default: float = 0.0) -> float:
    if val is None or val == "":
        return default
    try:
        s = str(val).replace(",", "").strip()
        return float(s)
    except (TypeError, ValueError):
        return default


def _parse_date(val: Any) -> datetime | None:
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.replace(tzinfo=UTC) if val.tzinfo is None else val
    s = str(val).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).replace(tzinfo=UTC)
        except ValueError:
            continue
    return None


def _write_rfq_lines_sheet(ws, *, business_type: str, include_sample: bool) -> None:
    ncols = len(RFQ_LINE_COLUMNS)
    ws["A1"] = TEMPLATE_VERSION
    ws["A1"].font = Font(bold=True, color="714B67")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws["A2"] = (
        "Ombi la Nukuu Bei / RFQ — Bei za kipimo ni bila kodi (unit prices ex-VAT). "
        "Kodi kwa kila mstari: none | vat_18. Images added later in app."
    )
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws["A3"] = f"Business type: {business_type} | Fill lines from row 6"
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)

    header_fill = PatternFill("solid", fgColor="F3F2F1")
    for col, (key, hint, _sample) in enumerate(RFQ_LINE_COLUMNS, start=1):
        ws.cell(row=4, column=col, value=hint)
        c1 = ws.cell(row=5, column=col, value=key)
        c1.font = Font(bold=True)
        c1.fill = header_fill
        if include_sample:
            ws.cell(row=6, column=col, value=_sample)


def _write_supplier_sheet(ws, *, include_sample: bool) -> None:
    ws["A1"] = "Msambazaji / Supplier (RFQ header)"
    ws["A1"].font = Font(bold=True, color="714B67")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    ws["A2"] = "Same fields as New RFQ / Purchase Request in the app. * = required for new vendor."
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
    headers = ["field_key", "label", "value", "required"]
    header_fill = PatternFill("solid", fgColor="F3F2F1")
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for i, (key, label, sample, req) in enumerate(SUPPLIER_SHEET_ROWS, start=5):
        ws.cell(row=i, column=1, value=key)
        ws.cell(row=i, column=2, value=label)
        if include_sample:
            ws.cell(row=i, column=3, value=sample)
        ws.cell(row=i, column=4, value=req)


def build_import_template_xlsx(*, business_type: str = "retail", include_sample: bool = True) -> bytes:
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl is required (pip install openpyxl)")

    wb = Workbook()
    ws_lines = wb.active
    ws_lines.title = "Agizo_Lines"
    _write_rfq_lines_sheet(ws_lines, business_type=business_type, include_sample=include_sample)

    ws_sup = wb.create_sheet("Msambazaji")
    _write_supplier_sheet(ws_sup, include_sample=include_sample)

    guide = wb.create_sheet("Guide")
    guide.append(["Topic", "Detail"])
    guide.append(["Unit prices", "Ex-VAT (bila kodi) — VAT computed per line when kodi = vat_18"])
    guide.append(["Supplier", "Fill Msambazaji sheet OR msambazaji column on each line"])
    guide.append(["Import", "Re-upload this file after editing to update inventory like PO receive"])
    guide.append(["Export", "Use Export inventory in app to download current stock in this format"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pick_lines_worksheet(wb):
    for name in ("Agizo_Lines", "RFQ_Lines", "Lines", "Products", "Bidhaa"):
        if name in wb.sheetnames:
            return wb[name]
    return wb.active


def _read_supplier_sheet(wb) -> dict[str, str]:
    for name in ("Msambazaji", "Supplier", "Msambazaji_RFQ"):
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        rows = [list(r) for r in ws.iter_rows(values_only=True) if r and any(c is not None and str(c).strip() for c in r)]
        out: dict[str, str] = {}
        for row in rows:
            cells = [str(c or "").strip() for c in row]
            if len(cells) >= 3 and cells[0].lower() in {k for k, *_ in SUPPLIER_SHEET_ROWS}:
                out[cells[0].lower()] = cells[2]
            elif len(cells) >= 2:
                key = _norm_header(cells[0]).replace(" ", "_")
                val = cells[1] if len(cells) == 2 else cells[-1]
                if key in {k for k, *_ in SUPPLIER_SHEET_ROWS} and val:
                    out[key] = str(val).strip()
        if out:
            return out
    return {}


def build_inventory_export_xlsx(
    products: list[dict[str, Any]],
    *,
    business_type: str = "retail",
    supplier_hint: dict[str, str] | None = None,
) -> bytes:
    """Export catalog in the same RFQ template shape for edit & re-import."""
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl is required (pip install openpyxl)")
    wb = Workbook()
    ws_lines = wb.active
    ws_lines.title = "Agizo_Lines"
    _write_rfq_lines_sheet(ws_lines, business_type=business_type, include_sample=False)
    start_row = 6
    for i, p in enumerate(products):
        meta = p.get("metadata_json") or p.get("metadata") or {}
        if not isinstance(meta, dict):
            meta = {}
        vat_type = str(meta.get("vat_type") or p.get("vat_type") or "").lower()
        line_tax = "vat_18" if _vat_type_is_taxable(vat_type) else "none"
        supplier = str(meta.get("supplier_name") or p.get("supplier") or "")
        row_vals = {
            "bidhaa": p.get("name"),
            "sku": p.get("sku"),
            "idadi": p.get("stock"),
            "bei_ya_kununua": p.get("cost"),
            "bei_ya_kuuza": p.get("price"),
            "kodi": line_tax,
            "kipimo": p.get("unit") or "pcs",
            "msambazaji": supplier,
            "barcode": p.get("barcode") or "",
            "category": p.get("category") or "",
            "batch_number": p.get("batch_number") or "",
            "expiry_date": (
                p.get("expiry_date").isoformat()[:10]
                if hasattr(p.get("expiry_date"), "isoformat")
                else (str(p.get("expiry_date") or "")[:10])
            ),
            "vat_type": vat_type or "standard",
            "reorder_level": p.get("reorder_point"),
            "shelf_location": meta.get("location") or "",
            "notes": meta.get("description") or "",
        }
        r = start_row + i
        for col, (key, _, __) in enumerate(RFQ_LINE_COLUMNS, start=1):
            ws_lines.cell(row=r, column=col, value=row_vals.get(key, ""))

    ws_sup = wb.create_sheet("Msambazaji")
    _write_supplier_sheet(ws_sup, include_sample=False)
    hint = supplier_hint or {}
    for i, (key, _label, _sample, _req) in enumerate(SUPPLIER_SHEET_ROWS, start=5):
        if key in hint and hint[key]:
            ws_sup.cell(row=i, column=3, value=hint[key])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_import_workbook(
    data: bytes,
    *,
    existing_skus: set[str],
    existing_barcodes: set[str],
    existing_by_sku: dict[str, dict[str, float]] | None = None,
    stock_mode: str = "absolute",
) -> dict[str, Any]:
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl is required (pip install openpyxl)")

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    supplier_info = _read_supplier_sheet(wb)
    ws = _pick_lines_worksheet(wb)
    rows_iter = ws.iter_rows(values_only=True)
    all_rows = [list(r) for r in rows_iter if r and any(c is not None and str(c).strip() for c in r)]
    if not all_rows:
        return {
            "template_detected": False,
            "rows": [],
            "summary": {"total": 0, "ready": 0, "error": 0},
            "supplier": supplier_info,
        }

    template_detected = any(TEMPLATE_VERSION in str(c) for row in all_rows[:6] for c in row if c)
    default_supplier_name = str(supplier_info.get("supplier_name") or "").strip()
    order_vat_scope = str(supplier_info.get("purchase_vat_scope") or "none").strip().lower()

    header_row_idx = 0
    col_map: dict[int, str] = {}
    for i, row in enumerate(all_rows[:30]):
        headers = [str(c or "") for c in row]
        trial = _map_headers(headers)
        if len(trial) >= 3:
            header_row_idx = i
            col_map = trial
            break
    if not col_map:
        return {
            "template_detected": template_detected,
            "rows": [],
            "summary": {"total": 0, "ready": 0, "error": 0, "warning": 0},
            "error": "Could not detect column headers. Download the Duka+ template or include name, sku, and price columns.",
        }

    out_rows: list[dict[str, Any]] = []
    seen_sku: set[str] = set()
    for i, row in enumerate(all_rows[header_row_idx + 1 :], start=header_row_idx + 2):
        mapped: dict[str, Any] = {}
        for col_idx, field in col_map.items():
            if col_idx < len(row):
                mapped[field] = row[col_idx]

        name = str(mapped.get("name") or "").strip()
        sku = str(mapped.get("sku") or "").strip()
        if not name and not sku:
            continue
        if not name:
            name = sku
        if not sku:
            sku = f"IMP-{secrets.token_hex(3).upper()}"

        raw_cost = _parse_float(mapped.get("cost"))
        price = _parse_float(mapped.get("price"))
        supplier = str(mapped.get("supplier") or "").strip() or default_supplier_name or "General Supplier"
        line_tax = _parse_line_tax(mapped.get("line_tax"))
        if line_tax == "none" and order_vat_scope == "all":
            line_tax = "vat_18"
        elif line_tax == "none" and order_vat_scope == "vat_products":
            vt_check = str(mapped.get("vat_type") or "").strip().lower()
            if _vat_type_is_taxable(vt_check):
                line_tax = "vat_18"
        vat_type_raw = str(mapped.get("vat_type") or "").strip().lower() or None
        cost_includes_vat = _parse_bool(mapped.get("cost_includes_vat")) if mapped.get("cost_includes_vat") not in (None, "") else None
        purchase_vat_rate = _parse_float(mapped.get("purchase_vat_rate"), default=-1.0)
        purchase_vat_rate = None if purchase_vat_rate < 0 else purchase_vat_rate

        messages: list[str] = []
        status = "ready"
        if raw_cost <= 0:
            if price > 0:
                raw_cost = round(price * 0.75, 2)
                messages.append("Cost missing — estimated at 75% of sell price")
                status = "warning"
            else:
                messages.append("Cost and price missing")
                status = "error"
        if price <= 0 and raw_cost > 0:
            price = round(raw_cost * 1.25, 2)
            messages.append("Sell price missing — set to cost + 25%")
            status = "warning" if status == "ready" else status
        if not str(mapped.get("supplier") or "").strip() and not default_supplier_name:
            messages.append("Supplier missing — fill Msambazaji sheet or msambazaji column")
            status = "warning" if status != "error" else status
        elif not str(mapped.get("supplier") or "").strip() and default_supplier_name:
            messages.append(f"Using supplier from Msambazaji sheet: {default_supplier_name}")

        unit_cost_ex_vat, cost_msgs = resolve_unit_cost_ex_vat(
            raw_cost,
            vat_type=vat_type_raw,
            cost_includes_vat=cost_includes_vat,
            purchase_vat_rate=purchase_vat_rate,
        )
        messages.extend(cost_msgs)
        stock_qty = _parse_float(mapped.get("stock"))

        sku_key = sku.lower()
        duplicate = "none"
        if sku_key in existing_skus:
            duplicate = "existing_sku"
            messages.append("SKU already in catalog — choose update or skip on import")
            status = "warning" if status != "error" else status
        elif sku_key in seen_sku:
            duplicate = "file_duplicate"
            messages.append("Duplicate SKU in file")
            status = "error"
        seen_sku.add(sku_key)

        barcode = str(mapped.get("barcode") or "").strip() or None
        if barcode and barcode in existing_barcodes:
            messages.append("Barcode already used")
            status = "warning" if status != "error" else status

        exp_dt = _parse_date(mapped.get("expiry_date"))
        prev = (existing_by_sku or {}).get(sku_key) or {}
        prev_stock = float(prev.get("stock") or 0)
        prev_cost = float(prev.get("cost") or 0)
        inventory_cost, stock_delta, new_stock = resolve_inventory_cost_after_import(
            prev_stock=prev_stock,
            prev_cost=prev_cost,
            target_stock=stock_qty,
            unit_cost_ex_vat=unit_cost_ex_vat,
            stock_mode=stock_mode,
        )
        if duplicate == "existing_sku" and stock_delta > 0 and prev_stock > 0:
            messages.append(
                f"Inventory cost blended (weighted avg): {inventory_cost:,.2f} TZS ex-VAT"
            )
        elif unit_cost_ex_vat > 0 and not cost_msgs:
            messages.append("Inventory unit cost ex-VAT (matches PO purchase line)")

        qty = stock_qty
        line_untaxed = round(qty * unit_cost_ex_vat, 2)
        tax_rate = DEFAULT_PURCHASE_VAT_RATE if line_tax == "vat_18" else 0.0
        line_tax_amount = round(line_untaxed * tax_rate, 2)
        product = {
            "name": name[:255],
            "sku": sku[:100],
            "barcode": barcode,
            "category": str(mapped.get("category") or "General").strip()[:100] or "General",
            "unit": str(mapped.get("unit") or "pcs").strip()[:30] or "pcs",
            "cost": inventory_cost if inventory_cost > 0 else round(unit_cost_ex_vat, 2),
            "unit_cost_ex_vat": round(unit_cost_ex_vat, 2),
            "raw_cost": round(raw_cost, 2),
            "price": round(price, 2),
            "stock": new_stock,
            "stock_delta": stock_delta,
            "reorder_point": _parse_float(mapped.get("reorder_point"), 10.0),
            "supplier": supplier,
            "batch_number": str(mapped.get("batch_number") or "").strip() or None,
            "expiry_date": exp_dt.date().isoformat() if exp_dt else None,
            "requires_prescription": _parse_bool(mapped.get("requires_prescription")),
            "location": str(mapped.get("location") or "").strip() or None,
            "description": str(mapped.get("description") or "").strip() or None,
            "vat_type": vat_type_raw,
            "cost_includes_vat": bool(cost_includes_vat),
            "line_tax": line_tax,
            "line_untaxed": line_untaxed,
            "line_tax_amount": line_tax_amount,
            "line_total": line_untaxed + line_tax_amount,
        }

        out_rows.append(
            {
                "row_number": i,
                "status": status,
                "messages": messages,
                "duplicate": duplicate,
                "product": product,
            }
        )

    ready = sum(1 for r in out_rows if r["status"] == "ready")
    warn = sum(1 for r in out_rows if r["status"] == "warning")
    err = sum(1 for r in out_rows if r["status"] == "error")
    return {
        "template_detected": template_detected or TEMPLATE_VERSION in str(all_rows[0][0] if all_rows else ""),
        "detected_fields": sorted(set(col_map.values())),
        "rows": out_rows,
        "summary": {"total": len(out_rows), "ready": ready, "warning": warn, "error": err},
        "supplier": supplier_info,
        "order_meta": {
            "purchase_vat_scope": order_vat_scope,
            "vat_note": supplier_info.get("vat_note"),
            "expected_date": supplier_info.get("expected_date"),
        },
    }


def _coerce_business_type(raw: str):
    from app.models import BusinessType

    try:
        return BusinessType(raw)
    except ValueError:
        return BusinessType.retail


async def _ensure_supplier_from_import(
    db,
    *,
    tenant_id: str,
    supplier_info: dict[str, Any] | None,
) -> int:
    """Create vendor record when missing (same fields as RFQ supplier form)."""
    if not supplier_info:
        return 0
    name = str(supplier_info.get("supplier_name") or "").strip()
    if not name:
        return 0
    try:
        from app.models import Supplier
    except Exception:
        return 0
    from sqlalchemy import select

    from sqlalchemy import func

    existing = (
        await db.execute(
            select(Supplier).where(
                Supplier.tenant_id == tenant_id,
                func.lower(Supplier.name) == name.lower(),
            )
        )
    ).scalar_one_or_none()
    if existing:
        return 0
    row = Supplier(
        tenant_id=tenant_id,
        name=name[:255],
        contact_person=str(supplier_info.get("contact_person") or "Sales")[:120],
        phone=str(supplier_info.get("phone") or "")[:40],
        email=str(supplier_info.get("email") or "")[:120],
        category=str(supplier_info.get("category") or "General")[:80],
        payment_terms=str(supplier_info.get("payment_terms") or "Net 30")[:80],
    )
    db.add(row)
    await db.flush()
    return 1


async def apply_product_import(
    db,
    *,
    tenant_id: str,
    branch_id: str | None,
    business_type: str,
    items: list[dict[str, Any]],
    duplicate_mode: str = "skip",
    stock_mode: str = "absolute",
    operator_name: str = "Excel import",
    supplier_info: dict[str, Any] | None = None,
) -> dict[str, Any]:
    from sqlalchemy import select

    from app.models import Product, StockMovement

    suppliers_created = await _ensure_supplier_from_import(db, tenant_id=tenant_id, supplier_info=supplier_info)

    existing = (
        await db.execute(select(Product).where(Product.tenant_id == tenant_id, Product.is_active == True))  # noqa: E712
    ).scalars().all()
    by_sku = {p.sku.lower(): p for p in existing}

    created = updated = skipped = failed = 0
    movements = 0
    errors: list[dict[str, str]] = []

    def _unit_cost_from_row(prod: dict[str, Any]) -> float:
        if prod.get("unit_cost_ex_vat") is not None:
            return float(prod.get("unit_cost_ex_vat") or 0)
        raw = float(prod.get("raw_cost") or prod.get("cost") or 0)
        incl = prod.get("cost_includes_vat")
        incl_bool = bool(incl) if incl is not None else None
        ex, _ = resolve_unit_cost_ex_vat(
            raw,
            vat_type=prod.get("vat_type"),
            cost_includes_vat=incl_bool,
            purchase_vat_rate=None,
        )
        return ex

    async def _log_stock_in(product: Product, qty: float, unit_cost: float, note: str) -> None:
        nonlocal movements
        if qty <= 0:
            return
        prev = float(product.stock or 0) - qty
        db.add(
            StockMovement(
                tenant_id=tenant_id,
                product_id=product.id,
                product_name=product.name,
                sku=product.sku,
                movement_type="in_purchase",
                quantity=qty,
                previous_stock=max(0.0, prev),
                new_stock=float(product.stock or 0),
                batch_number=product.batch_number,
                expiry_date=product.expiry_date,
                operator_name=operator_name,
                notes=note,
            )
        )
        movements += 1

    for item in items:
        prod = item.get("product") or item
        sku = str(prod.get("sku") or "").strip()
        if not sku:
            failed += 1
            continue
        meta: dict[str, Any] = {}
        if prod.get("description"):
            meta["description"] = prod["description"]
        if prod.get("location"):
            meta["location"] = prod["location"]
        if prod.get("supplier"):
            meta["supplier_name"] = prod["supplier"]
        if prod.get("vat_type"):
            meta["vat_type"] = prod["vat_type"]
        if prod.get("unit_cost_ex_vat") is not None:
            meta["last_purchase_unit_cost_ex_vat"] = prod["unit_cost_ex_vat"]
        if prod.get("line_tax"):
            meta["purchase_line_tax"] = prod["line_tax"]

        unit_ex = _unit_cost_from_row(prod)
        target_stock = float(prod.get("stock") or 0)

        existing_p = by_sku.get(sku.lower())
        if existing_p:
            if duplicate_mode == "skip":
                skipped += 1
                continue
            if duplicate_mode == "update":
                prev_stock = float(existing_p.stock or 0)
                prev_cost = float(existing_p.cost or 0)
                inv_cost, stock_delta, new_stock = resolve_inventory_cost_after_import(
                    prev_stock=prev_stock,
                    prev_cost=prev_cost,
                    target_stock=target_stock,
                    unit_cost_ex_vat=unit_ex,
                    stock_mode=stock_mode,
                )
                existing_p.name = prod.get("name") or existing_p.name
                existing_p.category = prod.get("category") or existing_p.category
                existing_p.price = float(prod.get("price") or existing_p.price)
                existing_p.cost = inv_cost if inv_cost > 0 else existing_p.cost
                existing_p.stock = new_stock
                existing_p.reorder_point = float(prod.get("reorder_point") or existing_p.reorder_point)
                existing_p.unit = prod.get("unit") or existing_p.unit
                existing_p.barcode = prod.get("barcode") or existing_p.barcode
                existing_p.batch_number = prod.get("batch_number") or existing_p.batch_number
                if prod.get("expiry_date"):
                    existing_p.expiry_date = _parse_date(prod.get("expiry_date")) or existing_p.expiry_date
                existing_p.requires_prescription = bool(prod.get("requires_prescription"))
                existing_p.metadata_json = {**(existing_p.metadata_json or {}), **meta}
                await db.flush()
                await _log_stock_in(
                    existing_p,
                    stock_delta,
                    unit_ex,
                    f"Excel import — stock-in (ex-VAT unit {unit_ex:,.2f})",
                )
                updated += 1
                continue
            # create_new
            sku = f"{sku}-IMP{secrets.token_hex(2).upper()}"

        inv_cost = unit_ex if unit_ex > 0 else float(prod.get("cost") or 0.01)
        if stock_mode == "add":
            new_stock = float(prod.get("stock") or 0)
            stock_delta = new_stock
        else:
            new_stock = float(prod.get("stock") or 0)
            stock_delta = new_stock

        row = Product(
            tenant_id=tenant_id,
            branch_id=branch_id,
            name=str(prod.get("name") or sku)[:255],
            category=str(prod.get("category") or "General")[:100],
            sku=sku[:100],
            barcode=prod.get("barcode"),
            price=float(prod.get("price") or 0),
            cost=inv_cost,
            stock=new_stock,
            reorder_point=float(prod.get("reorder_point") or 10),
            unit=str(prod.get("unit") or "pcs")[:30],
            batch_number=prod.get("batch_number"),
            expiry_date=_parse_date(prod.get("expiry_date")),
            requires_prescription=bool(prod.get("requires_prescription")),
            business_type=_coerce_business_type(business_type),
            metadata_json=meta,
        )
        db.add(row)
        await db.flush()
        by_sku[sku.lower()] = row
        await _log_stock_in(
            row,
            stock_delta,
            unit_ex,
            f"Excel import — opening stock (ex-VAT unit {unit_ex:,.2f})",
        )
        created += 1

    await db.flush()
    return {
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "failed": failed,
        "stock_movements": movements,
        "suppliers_created": suppliers_created,
        "errors": errors,
    }
