"""Zalongwa-style XLSX export from run_odoo_report payloads."""

from __future__ import annotations

import io
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def _header_row(ws, row: int, titles: list[str]) -> None:
    bold = Font(bold=True)
    for col, title in enumerate(titles, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = bold


def report_payload_to_xlsx_bytes(report_key: str, payload: dict[str, Any], company_name: str) -> bytes:
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl is required for XLSX export (pip install openpyxl)")

    wb = Workbook()
    ws = wb.active
    ws.title = report_key[:31]
    ws.cell(row=1, column=1, value=company_name)
    ws.cell(row=2, column=1, value=f"Report: {report_key}")
    ws.cell(row=3, column=1, value=f"Period: {payload.get('date_from')} — {payload.get('date_to')}")

    rtype = str(payload.get("type") or "")
    start = 5

    if rtype == "trial_balance":
        _header_row(ws, start, ["Code", "Account", "Initial", "Debit", "Credit", "End"])
        rows = payload.get("rows") or []
        for i, row in enumerate(rows, start=start + 1):
            ws.cell(row=i, column=1, value=row.get("code"))
            ws.cell(row=i, column=2, value=row.get("name"))
            ws.cell(row=i, column=3, value=row.get("initial_balance"))
            ws.cell(row=i, column=4, value=row.get("debit"))
            ws.cell(row=i, column=5, value=row.get("credit"))
            ws.cell(row=i, column=6, value=row.get("end_balance"))
        tr = payload.get("totals_row") or {}
        tot_row = start + len(rows) + 1
        if rows:
            td = tr.get("debit")
            tc = tr.get("credit")
            if td is None:
                td = round(sum(float(r.get("debit") or 0) for r in rows), 2)
            if tc is None:
                tc = round(sum(float(r.get("credit") or 0) for r in rows), 2)
            _header_row(ws, tot_row, ["", "TOTAL", "", td, tc, ""])
            ws.cell(row=tot_row + 1, column=2, value="Balanced" if abs(float(td) - float(tc)) < 0.02 else "Out of balance")

    elif rtype == "general_ledger":
        row_i = start
        sum_d = 0.0
        sum_c = 0.0
        for acct in payload.get("accounts") or []:
            sum_d += float(acct.get("debit") or 0)
            sum_c += float(acct.get("credit") or 0)
            ws.cell(row=row_i, column=1, value=f"{acct.get('code')} {acct.get('name')}")
            row_i += 1
            _header_row(ws, row_i, ["Date", "Reference", "Label", "Debit", "Credit"])
            row_i += 1
            for ln in acct.get("lines") or []:
                ws.cell(row=row_i, column=1, value=ln.get("date"))
                ws.cell(row=row_i, column=2, value=ln.get("reference"))
                ws.cell(row=row_i, column=3, value=ln.get("label"))
                ws.cell(row=row_i, column=4, value=ln.get("debit"))
                ws.cell(row=row_i, column=5, value=ln.get("credit"))
                row_i += 1
            row_i += 1
        if sum_d or sum_c:
            _header_row(ws, row_i, ["TOTAL GENERAL LEDGER", "", "", sum_d, sum_c])
            row_i += 1
            ws.cell(row=row_i, column=1, value="Balanced" if abs(sum_d - sum_c) < 0.02 else "Out of balance")
            row_i += 1

    elif rtype == "open_items":
        _header_row(ws, start, ["Date", "Move", "Partner", "Account", "Open amount"])
        for i, row in enumerate(payload.get("rows") or [], start=start + 1):
            ws.cell(row=i, column=1, value=row.get("date"))
            ws.cell(row=i, column=2, value=row.get("move_name"))
            ws.cell(row=i, column=3, value=row.get("partner_name"))
            ws.cell(row=i, column=4, value=row.get("account_code"))
            ws.cell(row=i, column=5, value=row.get("amount"))

    elif rtype in ("vat", "tax"):
        _header_row(ws, start, ["Description", "Amount"])
        for i, row in enumerate(payload.get("rows") or [], start=start + 1):
            ws.cell(row=i, column=1, value=row.get("label_en") or row.get("label"))
            ws.cell(row=i, column=2, value=row.get("amount"))

    elif rtype in ("hierarchy", "bs_odoo", "pnl_odoo"):
        _header_row(ws, start, ["Line", "Amount"])
        data = payload.get("data") or {}
        lines = payload.get("lines") or data.get("lines") or []
        for i, row in enumerate(lines, start=start + 1):
            indent = "  " * int(row.get("level") or 0)
            ws.cell(row=i, column=1, value=f"{indent}{row.get('name')}")
            ws.cell(row=i, column=2, value=row.get("amount"))

    elif rtype == "aged":
        _header_row(ws, start, ["Partner", "Not due", "1-30", "31-60", "61-90", "90+", "Total"])
        for i, row in enumerate(payload.get("rows") or [], start=start + 1):
            ws.cell(row=i, column=1, value=row.get("name"))
            ws.cell(row=i, column=2, value=row.get("current"))
            ws.cell(row=i, column=3, value=row.get("d30"))
            ws.cell(row=i, column=4, value=row.get("d60"))
            ws.cell(row=i, column=5, value=row.get("d90"))
            ws.cell(row=i, column=6, value=row.get("d90plus"))
            ws.cell(row=i, column=7, value=row.get("total"))

    elif rtype in ("tax", "vat", "fiscal"):
        _header_row(ws, start, ["Description", "Amount"])
        for i, row in enumerate(payload.get("rows") or [], start=start + 1):
            ws.cell(row=i, column=1, value=row.get("label_en") or row.get("label_sw") or row.get("label"))
            ws.cell(row=i, column=2, value=row.get("amount"))

    elif rtype in ("journal_report", "journal_audit"):
        row_i = start
        for entry in payload.get("entries") or []:
            ws.cell(row=row_i, column=1, value=f"{entry.get('date')} {entry.get('reference')}")
            row_i += 1
            _header_row(ws, row_i, ["Account", "Label", "Debit", "Credit"])
            row_i += 1
            for ln in entry.get("lines") or []:
                ws.cell(row=row_i, column=1, value=f"{ln.get('account_code')} {ln.get('account_name')}")
                ws.cell(row=row_i, column=2, value=ln.get("label"))
                ws.cell(row=row_i, column=3, value=ln.get("debit"))
                ws.cell(row=row_i, column=4, value=ln.get("credit"))
                row_i += 1
            row_i += 1

    elif rtype == "partner_ledger":
        row_i = start
        for key, title in (("customers", "Customers"), ("vendors", "Vendors")):
            rows = payload.get(key) or []
            if not rows:
                continue
            ws.cell(row=row_i, column=1, value=title)
            row_i += 1
            _header_row(ws, row_i, ["Partner", "Balance"])
            row_i += 1
            for r in rows:
                ws.cell(row=row_i, column=1, value=r.get("name"))
                ws.cell(row=row_i, column=2, value=r.get("total"))
                row_i += 1
            row_i += 1

    elif rtype == "executive_summary":
        _header_row(ws, start, ["KPI", "Amount"])
        for i, row in enumerate(payload.get("kpis") or [], start=start + 1):
            ws.cell(row=i, column=1, value=row.get("label_en") or row.get("label_sw"))
            ws.cell(row=i, column=2, value=row.get("amount"))

    elif rtype == "cash_flow":
        data = payload.get("data") or {}
        _header_row(ws, start, ["Line", "Amount"])
        for i, (label, key) in enumerate(
            [
                ("Cash opening", "cash_opening"),
                ("Operating", "operating"),
                ("Investing", "investing"),
                ("Financing", "financing"),
                ("Cash closing", "cash_closing"),
            ],
            start=start + 1,
        ):
            ws.cell(row=i, column=1, value=label)
            ws.cell(row=i, column=2, value=data.get(key))

    elif rtype == "budget_report":
        _header_row(ws, start, ["Line", "Actual", "Budget", "Variance"])
        for i, row in enumerate(payload.get("rows") or [], start=start + 1):
            ws.cell(row=i, column=1, value=row.get("label_en"))
            ws.cell(row=i, column=2, value=row.get("actual"))
            ws.cell(row=i, column=3, value=row.get("budget"))
            ws.cell(row=i, column=4, value=row.get("variance"))

    else:
        ws.cell(row=start, column=1, value=f"Export not implemented for report type: {rtype}")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
